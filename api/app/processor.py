import os
import shutil
import subprocess
import uuid
from pathlib import Path

import fitz
import openpyxl
from docx import Document
from PIL import Image

from .config import OUTPUT_DIR, TMP_DIR


def _make_temp_dir(suffix: str) -> Path:
    target = TMP_DIR / f"{uuid.uuid4()}-{suffix}"
    target.mkdir(parents=True, exist_ok=True)
    return target


def extract_pdf_text(pdf_path: str) -> str:
    doc = fitz.open(pdf_path)
    text_chunks = [page.get_text("text") for page in doc]
    doc.close()
    return "\n\n".join(chunk.strip() for chunk in text_chunks if chunk.strip())


def pdf_to_word(pdf_path: str, output_name: str) -> str:
    text = extract_pdf_text(pdf_path)
    doc = Document()
    doc.add_paragraph(text or "No text detected in the uploaded PDF.")
    output = OUTPUT_DIR / output_name
    doc.save(output)
    return str(output)


def pdf_to_text(pdf_path: str, output_name: str) -> str:
    text = extract_pdf_text(pdf_path)
    output = OUTPUT_DIR / output_name
    output.write_text(text or "No text detected in the uploaded PDF.", encoding="utf-8")
    return str(output)


def pdf_to_excel(pdf_path: str, output_name: str) -> str:
    text = extract_pdf_text(pdf_path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ExtractedText"
    rows = [line.strip() for line in text.splitlines() if line.strip()]
    for row_index, row_value in enumerate(rows, start=1):
        ws.cell(row=row_index, column=1, value=row_value)
    output = OUTPUT_DIR / output_name
    wb.save(output)
    return str(output)


def pdf_to_image(pdf_path: str, output_prefix: str) -> list[str]:
    doc = fitz.open(pdf_path)
    outputs = []
    for page_index in range(len(doc)):
        page = doc[page_index]
        image = page.get_pixmap(matrix=fitz.Matrix(2, 2))
        output = OUTPUT_DIR / f"{output_prefix}-page-{page_index + 1}.png"
        image.save(output)
        outputs.append(str(output))
    doc.close()
    return outputs


def image_to_pdf(image_path: str, output_name: str) -> str:
    image = Image.open(image_path)
    if image.mode not in ("RGB", "RGBA"):
        image = image.convert("RGB")
    output = OUTPUT_DIR / output_name
    image.save(output, "PDF", resolution=100.0)
    return str(output)


def word_to_pdf(word_path: str, output_name: str) -> str:
    output = OUTPUT_DIR / output_name
    command = [
        "soffice",
        "--headless",
        "--convert-to",
        "pdf:writer_pdf_Export",
        "--outdir",
        str(OUTPUT_DIR),
        word_path,
    ]
    completed = subprocess.run(command, capture_output=True, text=True)
    if completed.returncode != 0:
        raise RuntimeError(completed.stderr or "LibreOffice conversion failed")
    source_output = OUTPUT_DIR / Path(word_path).stem
    if source_output.with_suffix(".pdf").exists():
        pdf_path = source_output.with_suffix(".pdf")
    else:
        pdf_path = next(OUTPUT_DIR.glob(f"{Path(word_path).stem}*.pdf"), None)
    if pdf_path is None:
        raise FileNotFoundError("Converted PDF was not produced")
    if pdf_path.name != output_name:
        shutil.move(str(pdf_path), output)
    return str(output)


def merge_pdf(pdf_paths: list[str], output_name: str) -> str:
    merged = fitz.open()
    for pdf_path in pdf_paths:
        with fitz.open(pdf_path) as source:
            merged.insert_pdf(source)
    output = OUTPUT_DIR / output_name
    merged.save(output)
    merged.close()
    return str(output)


def split_pdf(pdf_path: str, output_prefix: str, page_ranges: list[str] | None = None) -> list[str]:
    source = fitz.open(pdf_path)
    output_files = []
    pages = page_ranges or [f"{index + 1}" for index in range(len(source))]
    for page_index, page_no in enumerate(pages):
        doc = fitz.open()
        doc.insert_pdf(source, from_page=int(page_no) - 1, to_page=int(page_no) - 1)
        output = OUTPUT_DIR / f"{output_prefix}-{page_index + 1}.pdf"
        doc.save(output)
        output_files.append(str(output))
    source.close()
    return output_files


def rotate_pdf(pdf_path: str, output_name: str, angle: int = 90) -> str:
    source = fitz.open(pdf_path)
    for page in source:
        page.set_rotation((page.rotation + angle) % 360)
    output = OUTPUT_DIR / output_name
    source.save(output)
    source.close()
    return str(output)


def compress_pdf(pdf_path: str, output_name: str) -> str:
    doc = fitz.open(pdf_path)
    output = OUTPUT_DIR / output_name
    doc.save(output, garbage=4, deflate=True)
    doc.close()
    return str(output)


def ocr_pdf(pdf_path: str, output_name: str) -> str:
    output = OUTPUT_DIR / output_name
    subprocess.run([
        "ocrmypdf",
        "--output-type",
        "pdf",
        str(pdf_path),
        str(output),
    ], check=True)
    return str(output)


def pdf_to_ocr_text(pdf_path: str, output_name: str) -> str:
    text = extract_pdf_text(pdf_path)
    output = OUTPUT_DIR / output_name
    output.write_text(text or "OCR text not detected.", encoding="utf-8")
    return str(output)
