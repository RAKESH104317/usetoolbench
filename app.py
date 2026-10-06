import io,os,re,uuid,sqlite3,zipfile,subprocess
from pathlib import Path
from datetime import datetime,timezone,timedelta
from fastapi import FastAPI,File,UploadFile,HTTPException,Form
from fastapi.responses import HTMLResponse,FileResponse
import fitz
import openpyxl
from docx import Document
from PIL import Image

BASE=Path(__file__).resolve().parent
ROOT=Path(os.getenv("STORAGE_ROOT",str(BASE/"storage"))); UP=ROOT/"uploads"; OUT=ROOT/"outputs"; DB=ROOT/"jobs.db"
MAX=int(os.getenv("MAX_FILE_SIZE_MB","25"))*1024*1024
RET=int(os.getenv("FILE_RETENTION_MINUTES","60"))
app=FastAPI(title="UseToolBench",version="2.0.0")
TOOLS=[("merge","Merge PDF",".pdf",1),("split","Split PDF",".pdf",0),("rotate","Rotate PDF",".pdf",0),("compress","Compress PDF",".pdf",0),("pdf-to-text","PDF to Text",".pdf",0),("pdf-to-word","PDF to Word",".pdf",0),("pdf-to-excel","PDF to Excel",".pdf",0),("pdf-to-image","PDF to Images",".pdf",0),("image-to-pdf","Image to PDF",".png,.jpg,.jpeg",1),("word-to-pdf","Word to PDF",".docx",0),("ocr","OCR PDF",".pdf",0)]

def init():
    for p in (UP,OUT):p.mkdir(parents=True,exist_ok=True)
    with sqlite3.connect(DB) as d:d.execute("CREATE TABLE IF NOT EXISTS jobs(id TEXT PRIMARY KEY,op TEXT,status TEXT,output TEXT,name TEXT,created TEXT,error TEXT)");d.commit()
    cutoff=datetime.now(timezone.utc)-timedelta(minutes=RET)
    with sqlite3.connect(DB) as d:
        for jid,p,_, in d.execute("SELECT id,output,created FROM jobs").fetchall():
            try:
                if datetime.fromisoformat(_)<cutoff:
                    if p:Path(p).unlink(missing_ok=True)
                    d.execute("DELETE FROM jobs WHERE id=?",(jid,))
            except:pass
        d.commit()
@app.on_event("startup")
def startup():init()
@app.get("/health")
def health():return {"status":"ok","service":"usetoolbench","version":"2.0.0"}
@app.get("/api/tools")
def tools():return {"tools":[{"id":a,"name":b,"accept":c,"multiple":bool(d)} for a,b,c,d in TOOLS],"max_file_size_mb":MAX//1048576}
def safe(n):return re.sub(r"[^A-Za-z0-9._-]+","-",Path(n).stem).strip(".-")[:70] or "document"
def newjob(op):
    j=str(uuid.uuid4())
    with sqlite3.connect(DB) as d:d.execute("INSERT INTO jobs VALUES(?,?,?,?,?,?,?)",(j,op,"processing",None,None,datetime.now(timezone.utc).isoformat(),None));d.commit()
    return j
def finish(j,p,n):
    with sqlite3.connect(DB) as d:d.execute("UPDATE jobs SET status='completed',output=?,name=? WHERE id=?",(str(p),n,j));d.commit()
def fail(j,e):
    with sqlite3.connect(DB) as d:d.execute("UPDATE jobs SET status='failed',error=? WHERE id=?",(str(e)[:1000],j));d.commit()
async def save(f,allowed):
    ext=Path(f.filename or "").suffix.lower()
    if ext not in allowed:raise HTTPException(400,"Unsupported file type.")
    p=UP/f"{uuid.uuid4()}{ext}";total=0
    try:
        with p.open("wb") as o:
            while c:=await f.read(1048576):
                total+=len(c)
                if total>MAX:raise HTTPException(413,f"Maximum file size is {MAX//1048576} MB.")
                o.write(c)
    except:p.unlink(missing_ok=True);raise
    if not total:p.unlink(missing_ok=True);raise HTTPException(400,"Empty file.")
    return p
def checkpdf(p):
    try:
        with p.open("rb") as f:
            if f.read(5)!=b"%PDF-":raise ValueError()
        d=fitz.open(p);ok=d.page_count>0;d.close()
        if not ok:raise ValueError()
    except:raise HTTPException(400,"Invalid PDF file.")
def txt(p):
    with fitz.open(p) as d:return "\n\n".join(x for x in (q.get_text("text").strip() for q in d) if x)
def dest(j,s):return OUT/(j+s)
def process(op,fs,j,pages,angle):
    b=safe(fs[0].name)
    if op=="merge":
        d=fitz.open()
        for p in fs:checkpdf(p);s=fitz.open(p);d.insert_pdf(s);s.close()
        q=dest(j,".pdf");d.save(q);d.close();return q,b+"-merged.pdf"
    p=fs[0];checkpdf(p) if op!="image-to-pdf" and op!="word-to-pdf" else None
    if op=="split":
        src=fitz.open(p);nums=[int(x)-1 for x in pages.split(",") if x.strip().isdigit()];nums=[x for x in nums if 0<=x<src.page_count]
        if not nums:src.close();raise HTTPException(400,"Enter valid pages such as 1,3,5.")
        d=fitz.open()
        for n in nums:d.insert_pdf(src,from_page=n,to_page=n)
        q=dest(j,".pdf");d.save(q);d.close();src.close();return q,b+"-split.pdf"
    if op=="rotate":
        a=int(angle)
        if a not in (90,180,270):raise HTTPException(400,"Angle must be 90, 180 or 270.")
        d=fitz.open(p)
        for x in d:x.set_rotation((x.rotation+a)%360)
        q=dest(j,".pdf");d.save(q);d.close();return q,b+"-rotated.pdf"
    if op=="compress":
        d=fitz.open(p);q=dest(j,".pdf");d.save(q,garbage=4,clean=True,deflate=True);d.close();return q,b+"-compressed.pdf"
    if op=="pdf-to-text":
        q=dest(j,".txt");q.write_text(txt(p) or "No selectable text found.",encoding="utf-8");return q,b+".txt"
    if op=="pdf-to-word":
        q=dest(j,".docx");d=Document();[d.add_paragraph(x) for x in (txt(p).split("\n\n") or ["No selectable text found."])];d.save(q);return q,b+".docx"
    if op=="pdf-to-excel":
        q=dest(j,".xlsx");w=openpyxl.Workbook();s=w.active
        for i,x in enumerate([x for x in txt(p).splitlines() if x.strip()],1):s.cell(i,1,x)
        w.save(q);return q,b+".xlsx"
    if op=="pdf-to-image":
        q=dest(j,".zip")
        with zipfile.ZipFile(q,"w",zipfile.ZIP_DEFLATED) as z:
            with fitz.open(p) as d:
                for i,x in enumerate(d,1):z.writestr(f"page-{i}.png",x.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False).tobytes("png"))
        return q,b+"-images.zip"
    if op=="image-to-pdf":
        q=dest(j,".pdf");d=fitz.open()
        for x in fs:
            im=Image.open(x).convert("RGB");m=io.BytesIO();im.save(m,"JPEG",quality=94);pg=d.new_page(width=im.width,height=im.height);pg.insert_image(pg.rect,stream=m.getvalue())
        d.save(q);d.close();return q,b+".pdf"
    if op=="word-to-pdf":
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import A4
        q=dest(j,".pdf");doc=Document(p);c=canvas.Canvas(str(q),pagesize=A4);y=A4[1]-50
        for para in doc.paragraphs:
            if para.text.strip():c.drawString(45,y,para.text[:120]);y-=16
            if y<45:c.showPage();y=A4[1]-50
        c.save();return q,b+".pdf"
    if op=="ocr":
        q=dest(j,".pdf")
        try:r=subprocess.run(["ocrmypdf","--skip-text","--deskew",str(p),str(q)],capture_output=True,text=True,timeout=300)
        except FileNotFoundError:raise HTTPException(503,"OCR is not installed on this server.")
        if r.returncode:raise RuntimeError(r.stderr[-700:] or "OCR failed.")
        return q,b+"-ocr.pdf"
    raise HTTPException(404,"Tool not found.")

@app.post("/api/process/{op}")
async def run(op:str,files:list[UploadFile]=File(...),pages:str=Form(""),angle:str=Form("90")):
    cfg={x[0]:x for x in TOOLS}.get(op)
    if not cfg:raise HTTPException(404,"Tool not found.")
    if not files:raise HTTPException(400,"Select a file.")
    if not cfg[3] and len(files)!=1:raise HTTPException(400,"This tool accepts one file.")
    if op=="merge" and len(files)<2:raise HTTPException(400,"Select at least two PDFs.")
    allowed=set(cfg[2].split(","));saved=[];j=newjob(op)
    try:
        for f in files:saved.append(await save(f,allowed))
        q,n=process(op,saved,j,pages,angle);finish(j,q,n)
        return {"job_id":j,"status":"completed","filename":n,"download_url":f"/api/jobs/{j}/download"}
    except Exception as e:fail(j,e);raise
    finally:
        for p in saved:p.unlink(missing_ok=True)
@app.get("/api/jobs/{j}")
def status(j):
    with sqlite3.connect(DB) as d:r=d.execute("SELECT * FROM jobs WHERE id=?",(j,)).fetchone()
    if not r:raise HTTPException(404,"Job not found.")
    return {"job_id":j,"operation":r[1],"status":r[2],"filename":r[4],"error":r[6],"download_url":f"/api/jobs/{j}/download" if r[2]=="completed" else None}
@app.get("/api/jobs/{j}/download")
def download(j):
    with sqlite3.connect(DB) as d:r=d.execute("SELECT * FROM jobs WHERE id=?",(j,)).fetchone()
    if not r or r[2]!="completed":raise HTTPException(404,"Output not found.")
    p=Path(r[3]).resolve()
    if p.parent!=OUT.resolve() or not p.is_file():raise HTTPException(404,"Output expired.")
    return FileResponse(p,filename=r[4],media_type="application/octet-stream",headers={"Cache-Control":"private,no-store"})
@app.get("/",response_class=HTMLResponse)
def home():return (BASE/"frontend"/"index.html").read_text(encoding="utf-8")
